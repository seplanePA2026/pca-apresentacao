# -*- coding: utf-8 -*-
from pathlib import Path
import json
import re
from datetime import datetime, date
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

FOLDER = Path(r'C:\Users\Lucas-PC\Documents\SEPLANE\PCAs_Enviados')
OUT = Path(r'C:\Users\Lucas-PC\Documents\SEPLANE\PCA\data\pcas_enviados.json')

FILE_ORG_MAP = {
    'CGM.xlsx': ('CGM', 'Controladoria Geral do Município'),
    'PGM.xlsx': ('PGM', 'Procuradoria Geral do Município'),
    'SDR.xlsx': ('SDR', 'Secretaria Municipal de Desenvolvimento Rural'),
    'SEAMM.xlsx': ('SEAMM', 'Secretaria Municipal de Agricultura, Meio Ambiente e Mineração'),
    'SECULTE.xlsx': ('SECULTE', 'Secretaria Municipal de Cultura e Turismo'),
    'SEDES.xlsx': ('SEDES', 'Secretaria Municipal de Desenvolvimento Social e Cidadania'),
    'SEFAZ.xlsx': ('SEFAZ', 'Secretaria Municipal de Fazenda'),
    'SEMMASP.xlsx': ('SEMMASP', 'Secretaria Municipal de Meio Ambiente e Serviços Públicos'),
    'SETIC.xlsx': ('SETIC', 'Secretaria Municipal de Tecnologia e Inovação'),
    'SME.xlsx': ('SME', 'Secretaria Municipal de Educação'),
    'SMOP.xlsx': ('SMOP', 'Secretaria Municipal de Obras e Projetos'),
    'SMS.xlsx': ('SMS', 'Secretaria Municipal de Saúde'),
    'GABIENTE VICE.docx': ('GVP', 'Gabinete do Vice-Prefeito'),
}

HEADER_ALIASES = {
    'dfd': ['n do dfd item', 'no do dfd item', 'nº do dfd item', 'dfd'],
    'org': ['secretaria requisitante', 'secretaria'],
    'sector': ['setor departamento', 'setor'],
    'type': ['tipo de objeto', 'tipo'],
    'description': ['descricao do item', 'descrição do item', 'descricao'],
    'company': ['empresa'],
    'renewal': ['renovacao de contrato', 'renovação de contrato'],
    'quantity': ['quantidade estimada', 'quantidade'],
    'unit': ['unidade de medida'],
    'unitValue': ['valor unitario estimado', 'valor unitário estimado'],
    'total': ['valor total estimado', 'valor total'],
    'priority': ['grau de prioridade', 'prioridade'],
    'period': ['periodo estimado', 'período estimado', 'trimestre estimado'],
    'linkage': ['vinculacao orcamentaria', 'vinculação orçamentária'],
}


def canon(s):
    s = str(s or '').strip().lower()
    s = (s.replace('á','a').replace('à','a').replace('â','a').replace('ã','a')
           .replace('é','e').replace('ê','e').replace('í','i')
           .replace('ó','o').replace('ô','o').replace('õ','o')
           .replace('ú','u').replace('ç','c'))
    s = re.sub(r'[^a-z0-9]+', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def to_number(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return 0.0
    s = s.replace('R$', '').replace(' ', '')
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def clean(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%m/%Y')
    if isinstance(v, date):
        return v.strftime('%m/%Y')
    return str(v).strip()


def is_footer(row_vals):
    joined = ' '.join(clean(x).lower() for x in row_vals if x is not None)
    return 'antes de enviar' in joined or joined.startswith('pca —') or 'modelo de preenchimento' in joined


def load_wb(path):
    try:
        return load_workbook(path, data_only=True, read_only=True)
    except Exception:
        # Fallback: strip broken drawings then retry via temp bytes in memory isn't easy;
        # try without data_only
        return load_workbook(path, data_only=False, read_only=True)


def parse_xlsx(path, code, default_name):
    wb = load_wb(path)
    # Prefer sheet named PCA
    sheet = None
    for sn in wb.sheetnames:
        if 'pca' in sn.lower():
            sheet = wb[sn]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=True))
    header_idx = None
    headers = []
    for i, row in enumerate(rows):
        cells = [clean(c) for c in row]
        if any('dfd' in canon(c) for c in cells) and any('secretaria' in canon(c) for c in cells):
            header_idx = i
            headers = [canon(c) for c in cells]
            break
    if header_idx is None:
        wb.close()
        return [], default_name, 'header not found'

    idx = {}
    for key, aliases in HEADER_ALIASES.items():
        for a in aliases:
            ca = canon(a)
            for hi, h in enumerate(headers):
                if h == ca or ca in h or h in ca:
                    idx[key] = hi
                    break
            if key in idx:
                break

    items = []
    org_name = default_name
    for row in rows[header_idx + 1:]:
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue
        if is_footer(row):
            continue
        desc = clean(row[idx['description']]) if 'description' in idx and idx['description'] < len(row) else ''
        sector = clean(row[idx['sector']]) if 'sector' in idx and idx['sector'] < len(row) else ''
        obj_type = clean(row[idx['type']]) if 'type' in idx and idx['type'] < len(row) else ''
        # skip incomplete placeholder rows
        if not desc and not sector and not obj_type:
            continue
        # require some substance
        qty = to_number(row[idx['quantity']]) if 'quantity' in idx and idx['quantity'] < len(row) else 0
        unit_value = to_number(row[idx['unitValue']]) if 'unitValue' in idx and idx['unitValue'] < len(row) else 0
        total = to_number(row[idx['total']]) if 'total' in idx and idx['total'] < len(row) else 0
        if total <= 0 and qty > 0 and unit_value > 0:
            total = qty * unit_value
        # skip empty value lines with no description
        if not desc:
            continue

        org_cell = clean(row[idx['org']]) if 'org' in idx and idx['org'] < len(row) else ''
        if org_cell:
            org_name = org_cell

        item = {
            'dfdNo': clean(row[idx['dfd']]) if 'dfd' in idx and idx['dfd'] < len(row) else '',
            'orgName': org_cell or org_name,
            'sector': sector or '—',
            'objectType': obj_type or '—',
            'description': desc,
            'company': clean(row[idx['company']]) if 'company' in idx and idx['company'] < len(row) else '',
            'contractRenewal': clean(row[idx['renewal']]) if 'renewal' in idx and idx['renewal'] < len(row) else '',
            'quantity': qty,
            'unitMeasure': clean(row[idx['unit']]) if 'unit' in idx and idx['unit'] < len(row) else '',
            'unitValue': unit_value,
            'total': round(total, 2),
            'priority': clean(row[idx['priority']]) if 'priority' in idx and idx['priority'] < len(row) else '',
            'period': clean(row[idx['period']]) if 'period' in idx and idx['period'] < len(row) else '',
            'budgetLink': clean(row[idx['linkage']]) if 'linkage' in idx and idx['linkage'] < len(row) else '',
        }
        items.append(item)

    wb.close()
    return items, org_name, None


def parse_docx(path, code, default_name):
    try:
        from docx import Document
    except ImportError:
        return [], default_name, 'python-docx missing'

    doc = Document(str(path))
    items = []
    org_name = default_name
    extra_aliases = {
        **HEADER_ALIASES,
        'type': HEADER_ALIASES['type'] + ['tipo do item', 'tipo'],
        'sector': HEADER_ALIASES['sector'] + ['unidade demandante', 'unidade'],
        'description': HEADER_ALIASES['description'] + ['descricao da pretensao contratual', 'pretensao contratual'],
        'period': HEADER_ALIASES['period'] + ['data desejada'],
        'total': HEADER_ALIASES['total'] + ['valor estimado anual', 'valor estimado'],
    }

    for table in doc.tables:
        rows = [[clean(c.text) for c in row.cells] for row in table.rows]
        if not rows:
            continue
        header_row = 0
        headers = [canon(h) for h in rows[0]]
        if not any('descricao' in h or 'valor' in h or 'tipo' in h for h in headers):
            continue
        data_rows = rows[header_row + 1:]

        idx = {}
        for key, aliases in extra_aliases.items():
            for a in aliases:
                ca = canon(a)
                for hi, h in enumerate(headers):
                    if h == ca or ca in h or h in ca:
                        idx[key] = hi
                        break
                if key in idx:
                    break
        if 'description' not in idx and 'total' not in idx:
            continue

        for row in data_rows:
            if not any(row):
                continue
            desc = row[idx['description']] if 'description' in idx and idx['description'] < len(row) else ''
            if not desc:
                continue
            qty = to_number(row[idx['quantity']]) if 'quantity' in idx and idx['quantity'] < len(row) else 0
            unit_value = to_number(row[idx['unitValue']]) if 'unitValue' in idx and idx['unitValue'] < len(row) else 0
            total = to_number(row[idx['total']]) if 'total' in idx and idx['total'] < len(row) else 0
            if total <= 0 and qty > 0 and unit_value > 0:
                total = qty * unit_value
            sector = row[idx['sector']] if 'sector' in idx and idx['sector'] < len(row) else code
            items.append({
                'dfdNo': row[idx['dfd']] if 'dfd' in idx and idx['dfd'] < len(row) else '',
                'orgName': org_name,
                'sector': sector or code,
                'objectType': row[idx['type']] if 'type' in idx and idx['type'] < len(row) else '—',
                'description': desc,
                'company': row[idx['company']] if 'company' in idx and idx['company'] < len(row) else '',
                'contractRenewal': row[idx['renewal']] if 'renewal' in idx and idx['renewal'] < len(row) else '',
                'quantity': qty or 1,
                'unitMeasure': row[idx['unit']] if 'unit' in idx and idx['unit'] < len(row) else 'ano',
                'unitValue': unit_value or total,
                'total': round(total, 2),
                'priority': row[idx['priority']] if 'priority' in idx and idx['priority'] < len(row) else '',
                'period': row[idx['period']] if 'period' in idx and idx['period'] < len(row) else '',
                'budgetLink': row[idx['linkage']] if 'linkage' in idx and idx['linkage'] < len(row) else '',
            })
    return items, org_name, None if items else 'no table rows'


def main():
    organs = []
    errors = []
    for fname, (code, default_name) in FILE_ORG_MAP.items():
        path = FOLDER / fname
        if not path.exists():
            errors.append(f'{fname}: missing')
            continue
        try:
            if path.suffix.lower() == '.xlsx':
                items, org_name, err = parse_xlsx(path, code, default_name)
            else:
                items, org_name, err = parse_docx(path, code, default_name)
            if err:
                errors.append(f'{fname}: {err}')
            total = round(sum(i['total'] for i in items), 2)
            # Prefer spreadsheet name when present; keep file acronym as code
            display_name = org_name or default_name
            sectors = {}
            for it in items:
                sectors.setdefault(it['sector'] or '—', 0)
                sectors[it['sector'] or '—'] += it['total']
                it['orgCode'] = code
                it['orgName'] = display_name
            organs.append({
                'code': code,
                'name': display_name,
                'file': fname,
                'itemCount': len(items),
                'total': total,
                'sectors': [
                    {
                        'name': sname,
                        'total': round(stotal, 2),
                        'itemCount': sum(1 for i in items if (i['sector'] or '—') == sname),
                        'items': [i for i in items if (i['sector'] or '—') == sname],
                    }
                    for sname, stotal in sorted(sectors.items(), key=lambda x: (-x[1], x[0]))
                ],
                'items': items,
            })
            print(f'{code:8} {len(items):4} itens  R$ {total:,.2f}  | {org_name}')
        except Exception as e:
            errors.append(f'{fname}: {type(e).__name__}: {e}')
            print(f'ERROR {fname}: {e}')

    organs.sort(key=lambda o: o['name'])
    payload = {
        'year': 2027,
        'source': 'PCAs_Enviados',
        'generatedAt': datetime.now().isoformat(timespec='seconds'),
        'organCount': len(organs),
        'itemCount': sum(o['itemCount'] for o in organs),
        'total': round(sum(o['total'] for o in organs), 2),
        'organs': organs,
        'errors': errors,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print('---')
    print(f'Organs: {payload["organCount"]} | Items: {payload["itemCount"]} | Total: R$ {payload["total"]:,.2f}')
    print('Wrote', OUT)
    if errors:
        print('Errors:')
        for e in errors:
            print(' -', e)


if __name__ == '__main__':
    main()

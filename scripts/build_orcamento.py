# -*- coding: utf-8 -*-
"""Extrai Apresentação (1).xlsx → data/orcamento.json (aba Luquinha)."""
from pathlib import Path
import json
from datetime import datetime
import openpyxl

SRC = Path(r'C:\Users\Lucas-PC\Documents\SEPLANE\PCA')
files = [p for p in SRC.iterdir() if p.suffix.lower() == '.xlsx' and 'apresenta' in p.name.lower()]
if not files:
    raise SystemExit('Apresentacao xlsx not found')
path = files[0]
OUT = SRC / 'data' / 'orcamento.json'

wb = openpyxl.load_workbook(path, data_only=True)
formula_sheet = openpyxl.load_workbook(path, data_only=False)['Luquinha']
ws = wb['Luquinha'] if 'Luquinha' in wb.sheetnames else wb[wb.sheetnames[0]]

# "Gabinete" is a header that equals the sum of Manut.Gab / Centro / Ouvidoria / Assessoria
PARENT_HEADERS = {'gabinete'}

rows = []
for r in range(4, 27):
    name = ws.cell(r, 1).value
    if not name:
        continue
    name = str(name).strip()
    if name.lower() == 'total':
        continue
    if name.lower().strip() in PARENT_HEADERS:
        continue

    initial = float(ws.cell(r, 2).value or 0)
    pct2026 = float(ws.cell(r, 3).value or 0)
    pct2027 = float(ws.cell(r, 4).value or 0)
    updated = float(ws.cell(r, 5).value or 0)
    pctUpdated = float(ws.cell(r, 6).value or 0)
    paid = float(ws.cell(r, 7).value or 0)
    forecastDec = float(ws.cell(r, 8).value or 0)
    forecast2027 = float(ws.cell(r, 9).value or 0)
    delta = forecast2027 - initial
    deltaPct = (delta / initial * 100) if initial else None
    shareDelta = pct2027 - pct2026
    rows.append({
        'name': name,
        'initial2026': round(initial, 2),
        'share2026': round(pct2026, 6),
        'share2027': round(pct2027, 6),
        'updated': round(updated, 2),
        'shareUpdated': round(pctUpdated, 6),
        'paid': round(paid, 2),
        'forecastDec2026': round(forecastDec, 2),
        'forecast2027': round(forecast2027, 2),
        'deltaValue': round(delta, 2),
        'deltaPct': None if deltaPct is None else round(deltaPct, 4),
        'shareDeltaPp': round(shareDelta, 6),
    })

# Official totals from sheet Total row
total_initial = float(ws.cell(28, 2).value or sum(r['initial2026'] for r in rows))
total_updated = float(ws.cell(28, 5).value or sum(r['updated'] for r in rows))
total_paid = float(ws.cell(28, 7).value or sum(r['paid'] for r in rows))
total_forecast_dec = float(ws.cell(28, 8).value or sum(r['forecastDec2026'] for r in rows))
total_forecast_2027 = float(ws.cell(28, 9).value or sum(r['forecast2027'] for r in rows))

# Recompute shares against official total so they sum ~100%
for r in rows:
    r['share2026'] = round((r['initial2026'] / total_initial) * 100, 6) if total_initial else 0
    r['share2027'] = round((r['forecast2027'] / total_forecast_2027) * 100, 6) if total_forecast_2027 else 0
    r['shareUpdated'] = round((r['updated'] / total_updated) * 100, 6) if total_updated else 0
    r['shareDeltaPp'] = round(r['share2027'] - r['share2026'], 6)

revenue = {
    'planned2026': float(ws.cell(31, 2).value or total_initial),
    'collectedSep9': float(ws.cell(33, 7).value or 0),
    'forecastDec2026': round(float(ws.cell(33, 8).value or 0), 2),
    'forecast2027': float(ws.cell(33, 9).value or 0),
    'balanceSep9': round(float(ws.cell(37, 7).value or 0), 2),
    'balanceForecastDec2026': round(float(ws.cell(37, 8).value or 0), 2),
}

payload = {
    'source': path.name,
    'sheet': ws.title,
    'referenceDate': '2026-09-09',
    'generatedAt': datetime.now().isoformat(timespec='seconds'),
    'totals': {
        'initial2026': round(total_initial, 2),
        'updated': round(total_updated, 2),
        'paid': round(total_paid, 2),
        'forecastDec2026': round(total_forecast_dec, 2),
        'forecast2027': round(total_forecast_2027, 2),
        'deltaValue': round(total_forecast_2027 - total_initial, 2),
        'deltaPct': round((total_forecast_2027 - total_initial) / total_initial * 100, 4) if total_initial else 0,
        'executionPct': round(total_paid / total_updated * 100, 4) if total_updated else 0,
    },
    'revenue': revenue,
    'organs': rows,
    'formulaRules': [
        {
            'fixedForecast': not str(formula_sheet.cell(r, 9).value).startswith('='),
            'fixedShare': formula_sheet.cell(r, 6).value if isinstance(formula_sheet.cell(r, 6).value, (int, float)) else None,
        }
        for r in range(4, 27) if r != 5
    ],
}

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
print('Wrote', OUT)
print('Organs', len(rows))
print('2026', total_initial)
print('2027', total_forecast_2027)
print('Delta', total_forecast_2027 - total_initial)
